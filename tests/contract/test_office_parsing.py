import hashlib
import io
import sys
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "data-pipeline/src"))

import docx
from pptx import Presentation
from pptx.util import Inches

from biz_aid_pipeline.chunking.chunker import item_provenance
from biz_aid_pipeline.parsing.models import ParseRequest, parse_identity, parsing_contract
from biz_aid_pipeline.parsing.router import parse_document


def docx_bytes(empty=False):
    document = docx.Document()
    if not empty:
        document.add_heading("2026년 소상공인 지원사업 모집 공고", level=1)
        document.add_heading("1. 지원 대상", level=2)
        document.add_paragraph("서울시 소재 소상공인으로 사업자등록 후 6개월이 지난 기업을 대상으로 합니다.")
        table = document.add_table(rows=3, cols=2)
        for row, values in enumerate((("구분", "내용"), ("지원 규모", "기업당 최대 500만원"), ("신청 기간", "2026년 12월까지"))):
            for col, value in enumerate(values):
                table.cell(row, col).text = value
        document.add_heading("2. 신청 방법", level=2)
        document.add_paragraph("온라인 접수 시스템에서 신청서를 제출합니다.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pptx_bytes():
    presentation = Presentation()
    for title, body in (("상품기술서", "지역상생협력관 코칭 상담 신청을 위한 상품기술서 양식입니다."),
                        ("상품 정보", "상품명과 구성, 가격, 유통 경로를 적어 주세요. 예시를 지우고 작성합니다.")):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    table = presentation.slides[1].shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
    for (row, col), value in {(0, 0): "항목", (0, 1): "내용", (1, 0): "상품명", (1, 1): "맛있는 스낵바"}.items():
        table.cell(row, col).text = value
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def no_libreoffice():
    """Docling backend가 외부 LibreOffice 변환기를 찾으면 실패시킨다(backend 모듈이 import한 이름을 막는다)."""
    guard = AssertionError("LibreOffice converter requested")
    return mock.patch.multiple("docling.backend.msword_backend", get_docx_to_pdf_converter=mock.Mock(side_effect=guard)), \
        mock.patch.multiple("docling.backend.mspowerpoint_backend", get_docx_to_pdf_converter=mock.Mock(side_effect=guard))


def parse(raw, detected):
    return parse_document(ParseRequest(hashlib.sha256(raw).hexdigest(), detected, len(raw)), raw, parsing_contract())


def office_meta(item):
    return item.meta.get_custom_part()["bizaid__office"]


class OfficeParsingTests(unittest.TestCase):
    def test_docx_keeps_order_heading_path_and_table_without_fake_pages(self):
        # BOUNDARY: 설치 여부와 무관하게 외부 LibreOffice 변환을 부르지 않는다(부르면 test가 실패한다).
        word, powerpoint = no_libreoffice()
        with word, powerpoint:
            result = parse(docx_bytes(), "DOCX")
        self.assertEqual((result.route, result.status, result.failure_code), ("DOCLING_DOCX", "PARSED", None))
        document = result.document
        self.assertEqual(len(document.pages), 0)
        paragraph = next(item for item in document.texts if item.text.startswith("서울시 소재"))
        meta = office_meta(paragraph)
        self.assertEqual(meta["heading_path"], ["2026년 소상공인 지원사업 모집 공고", "1. 지원 대상"])
        self.assertEqual((meta["format"], meta["source_sha256"]), ("DOCX", result.source_sha256))
        later = next(item for item in document.texts if item.text.startswith("온라인 접수"))
        self.assertEqual(office_meta(later)["heading_path"], ["2026년 소상공인 지원사업 모집 공고", "2. 신청 방법"])
        self.assertLess(meta["order"], office_meta(later)["order"])
        # 표는 행·열 구조와 위치 meta를 함께 가진다.
        table = document.tables[0]
        self.assertEqual((table.data.num_rows, table.data.num_cols), (3, 2))
        self.assertEqual([cell.text for cell in table.data.grid[1]], ["지원 규모", "기업당 최대 500만원"])
        self.assertEqual(office_meta(table)["heading_path"][-1], "1. 지원 대상")
        # chunk provenance는 page·bbox 없이 문서 순서와 제목 경로를 옮긴다.
        entry = item_provenance(paragraph, document)
        self.assertNotIn("page", entry)
        self.assertEqual((entry["block_order"], entry["heading_path"][-1]), (meta["order"], "1. 지원 대상"))
        self.assertEqual(document.name, result.source_sha256)

    def test_pptx_keeps_slide_numbers_and_tables(self):
        word, powerpoint = no_libreoffice()
        with word, powerpoint:
            result = parse(pptx_bytes(), "PPTX")
        self.assertEqual((result.route, result.status, result.unit_count), ("DOCLING_PPTX", "PARSED", 2))
        document = result.document
        body = next(item for item in document.texts if item.text.startswith("상품명과 구성"))
        self.assertEqual(office_meta(body)["slide"], 2)
        entry = item_provenance(body, document)
        self.assertEqual((entry["slide"], entry["page"]), (2, 2))
        # 위치 단위는 다른 형식과 같은 pt다(기본 4:3 슬라이드 = 10 × 7.5 inch = 720 × 540 pt). EMU 값이 남으면 수십만 단위가 된다.
        self.assertEqual({(page.size.width, page.size.height) for page in document.pages.values()}, {(720.0, 540.0)})
        left, top, right, bottom = entry["bbox_pt"]
        self.assertTrue(0 <= left < right <= 720 and 0 <= top < bottom <= 540, entry["bbox_pt"])
        table = document.tables[0]
        self.assertEqual((table.data.num_rows, table.data.num_cols, office_meta(table)["slide"]), (2, 2, 2))

    def test_empty_document_and_broken_bytes(self):
        self.assertEqual(parse(docx_bytes(empty=True), "DOCX").status, "EMPTY_TEXT")
        broken = parse(b"PK\x03\x04 not a real office file", "DOCX")
        self.assertEqual(broken.status, "PARSE_FAILED")
        self.assertTrue(broken.failure_code.startswith("office_"))

    def test_parse_key_scope_has_office_inputs_only(self):
        contract = parsing_contract()
        for route, library in (("DOCLING_DOCX", "python_docx_version"), ("DOCLING_PPTX", "python_pptx_version")):
            identity = parse_identity("0" * 64, route, contract)
            present = {key for key, value in identity.items() if value is not None}
            self.assertEqual(present, {"source_sha256", "route", "normalizer_version", "docling_core_version", "docling_version",
                                       "office_parser_version", "office_config_sha256", library})
        # 근거 위치(page·bbox·provenance)는 어떤 식별값의 입력도 아니다. 단위를 바꿔도 key가 바뀌지 않는다.
        from biz_aid_pipeline.chunking.chunker import chunking_contract
        from biz_aid_pipeline.indexing.embedder import indexing_contract
        chunk_inputs = chunking_contract()["identity"]["chunk_set_key_inputs"]
        embedding_inputs = indexing_contract()["identity"]["embedding_key_inputs"]
        for name in ("provenance", "bbox_pt", "page", "pages", "slide"):
            self.assertNotIn(name, chunk_inputs)
            self.assertNotIn(name, embedding_inputs)
            self.assertNotIn(name, identity)
        # 기존 route identity에는 Office 전용 key가 생기지 않는다.
        self.assertNotIn("office_parser_version", parse_identity("0" * 64, "DOCLING_PDF", contract))


if __name__ == "__main__":
    unittest.main()
